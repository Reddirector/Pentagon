from __future__ import annotations

import asyncio
import io
import logging
from pathlib import PurePath
from uuid import uuid4

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Conversation, Document, User
from app.db.session import get_db
from app.schemas import DocumentResponse
from app.security.keys import resolve_api_key_or_http
from app.services.document_store import delete_document_chunks, store_chunks


logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/documents", tags=["documents"])
_MAX_UPLOAD_BYTES = 100 * 1024 * 1024
_SUPPORTED_EXTENSIONS = {
    # Documents
    ".pdf", ".docx", ".doc", ".rtf", ".odt",
    # Presentations
    ".pptx", ".ppt",
    # Spreadsheets
    ".xlsx", ".xls", ".csv", ".tsv",
    # Markup & Web
    ".md", ".markdown", ".mdown", ".html", ".htm", ".xhtml", ".xml", ".rst",
    # Structured data & configs
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".conf", ".env",
    # Code & scripts
    ".txt", ".log", ".py", ".js", ".jsx", ".ts", ".tsx", ".sh", ".bash",
    ".sql", ".css", ".scss", ".c", ".cpp", ".h", ".hpp", ".go", ".rs",
    ".java", ".rb", ".php", ".diff", ".patch",
}


@router.post("/upload", response_model=DocumentResponse, status_code=201)
async def upload_document(
    file: UploadFile = File(...),
    user_id: str = Form(min_length=1, max_length=128),
    conversation_id: str = Form(min_length=1, max_length=36),
    db: Session = Depends(get_db),
) -> dict:
    conversation = db.scalar(
        select(Conversation).where(
            Conversation.id == conversation_id,
            Conversation.user_id == user_id,
        )
    )
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found for this user.")

    filename = _safe_filename(file.filename or "")
    extension = PurePath(filename).suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file format. Supported: PDF, Word, Excel, PowerPoint, CSV, Markdown, HTML, JSON, YAML, and text/code files.",
        )

    raw_file = await file.read(_MAX_UPLOAD_BYTES + 1)
    if len(raw_file) > _MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The maximum document size is 100 MB.")
    try:
        text = await asyncio.to_thread(_extract_text, raw_file, extension)
    except Exception as exc:
        logger.info("Document extraction failed (%s)", type(exc).__name__)
        raise HTTPException(status_code=422, detail="Could not extract text from this document.") from None
    if not text.strip():
        raise HTTPException(status_code=422, detail="This document contains no extractable text.")

    chunks = await asyncio.to_thread(_chunk_text, text)
    if not chunks:
        raise HTTPException(status_code=422, detail="This document contains no extractable text.")

    try:
        api_key: str | None = resolve_api_key_or_http(db, user_id)
    except HTTPException as exc:
        # Unlike chat, an upload can still be indexed without a user key: the
        # server-side fallback (or a keyless embedding provider) may cover it.
        if exc.status_code != 404:
            raise
        if db.get(User, user_id) is None:
            raise HTTPException(status_code=404, detail="User not found. Store an API key first.") from None
        api_key = None

    document_id = str(uuid4())
    try:
        collection_name, chunks_stored = await store_chunks(
            conversation_id=conversation_id,
            document_id=document_id,
            filename=filename,
            chunks=chunks,
            user_id=user_id,
            api_key=api_key,
        )
    except Exception as exc:
        logger.warning("Document vector storage failed (%s)", type(exc).__name__)
        raise HTTPException(status_code=502, detail="Could not embed or store this document.") from None

    document = Document(
        id=document_id,
        user_id=user_id,
        conversation_id=conversation_id,
        collection_name=collection_name,
        filename=filename,
        chunk_count=chunks_stored,
    )
    db.add(document)
    try:
        db.commit()
    except Exception:
        db.rollback()
        try:
            await delete_document_chunks(collection_name, document_id)
        except Exception:
            logger.warning("Orphan document chunks could not be removed after database failure")
        raise HTTPException(status_code=500, detail="Could not save document metadata.") from None

    return {
        "document_id": document.id,
        "filename": document.filename,
        "chunks_stored": document.chunk_count,
        "created_at": document.created_at,
    }


@router.delete("/{document_id}")
async def remove_document(
    document_id: str,
    user_id: str = Query(min_length=1, max_length=128),
    db: Session = Depends(get_db),
) -> dict[str, object]:
    document = db.scalar(
        select(Document).where(Document.id == document_id, Document.user_id == user_id)
    )
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found for this user.")

    try:
        await delete_document_chunks(document.collection_name, document.id)
    except Exception as exc:
        logger.warning("Document vector deletion failed (%s)", type(exc).__name__)
        raise HTTPException(status_code=502, detail="Could not remove document chunks.") from None
    db.delete(document)
    db.commit()
    return {"deleted": True, "document_id": document_id}


@router.get("")
def list_documents(
    user_id: str = Query(min_length=1, max_length=128),
    conversation_id: str = Query(min_length=1, max_length=36),
    db: Session = Depends(get_db),
) -> list[dict[str, object]]:
    conversation = db.scalar(
        select(Conversation).where(
            Conversation.id == conversation_id,
            Conversation.user_id == user_id,
        )
    )
    if conversation is None:
        raise HTTPException(status_code=404, detail="Conversation not found for this user.")
    rows = db.scalars(
        select(Document)
        .where(Document.user_id == user_id, Document.conversation_id == conversation_id)
        .order_by(Document.created_at)
    )
    return [
        {
            "document_id": item.id,
            "filename": item.filename,
            "chunks_stored": item.chunk_count,
            "created_at": item.created_at,
        }
        for item in rows
    ]


def _safe_filename(filename: str) -> str:
    normalized = filename.replace("\\", "/").split("/")[-1].strip()
    return normalized[:255]


def _decode_bytes(raw_file: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return raw_file.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw_file.decode("utf-8", errors="replace")


def _extract_text(raw_file: bytes, extension: str) -> str:
    import csv
    import re

    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(raw_file), strict=False)
        pages_out = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages_out.append(f"--- Page {index} ---\n{page_text.strip()}")
            except Exception as exc:
                logger.debug("PDF page %d extract error: %s", index, exc)
        return "\n\n".join(pages_out) if pages_out else "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        document = WordDocument(io.BytesIO(raw_file))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            )
        return "\n".join(paragraphs)

    if extension == ".doc":
        try:
            document = WordDocument(io.BytesIO(raw_file))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            if paragraphs:
                return "\n".join(paragraphs)
        except Exception:
            pass
        return _decode_bytes(raw_file)

    if extension in {".pptx", ".ppt"}:
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(raw_file))
            slides_out = []
            for i, slide in enumerate(prs.slides, start=1):
                slide_content = [f"--- Slide {i} ---"]
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text:
                        txt = shape.text_frame.text.strip()
                        if txt:
                            slide_content.append(txt)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text.strip():
                                slide_content.append(row_text)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_content.append(f"Notes: {notes}")
                if len(slide_content) > 1:
                    slides_out.append("\n".join(slide_content))
            return "\n\n".join(slides_out)
        except Exception as exc:
            logger.warning("PowerPoint extraction failed: %s", exc)
            return _decode_bytes(raw_file)

    if extension in {".xlsx", ".xls"}:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(raw_file), data_only=True, read_only=True)
            sheets_out = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None and str(c).strip() for c in row):
                        rows.append(" | ".join("" if c is None else str(c).strip() for c in row))
                if rows:
                    sheets_out.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
            return "\n\n".join(sheets_out)
        except Exception as exc:
            logger.warning("Spreadsheet extraction failed: %s", exc)
            return _decode_bytes(raw_file)

    if extension in {".csv", ".tsv"}:
        text = _decode_bytes(raw_file)
        delimiter = "\t" if extension == ".tsv" else ","
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            rows = [" | ".join(cell.strip() for cell in row) for row in reader if any(cell.strip() for cell in row)]
            return "\n".join(rows) if rows else text
        except Exception:
            return text

    if extension in {".html", ".htm", ".xhtml"}:
        text = _decode_bytes(raw_file)
        try:
            import trafilatura

            extracted = trafilatura.extract(text)
            if extracted and extracted.strip():
                return extracted.strip()
        except Exception:
            pass
        clean = re.sub(r"<script.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<style.*?</style>", "", clean, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<[^>]+>", " ", clean)
        return re.sub(r"\s+", " ", clean).strip()

    return _decode_bytes(raw_file)


def _chunk_text(text: str) -> list[str]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=2000,
        chunk_overlap=200,
        length_function=len,
    )
    return [chunk.strip() for chunk in splitter.split_text(text) if chunk.strip()]
