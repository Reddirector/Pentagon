"""Generated files: generation, serving, ownership, deletion, and the gate.

The five acceptance questions this file answers:

* does each format produce a real, correctly-typed file;
* do the routes serve metadata, byte-identical downloads and a preview;
* is a file from conversation A invisible to conversation B and to other
  users (checked against the API, not the UI);
* does deleting from the management list remove the object *and* the row;
* does writing the file actually stop at the permission gate.
"""

from __future__ import annotations

import asyncio
from uuid import uuid4

import httpx
import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.types import Command

from app.agent.schemas import ToolContext
from app.config import settings
from app.db.models import ActionAuditLog, AutonomySetting, Conversation, GeneratedFile, User
from app.db.session import SessionLocal, initialize_database
from app.main import app
from app.services import chat_graph, file_storage
from app.services.file_storage import LocalStorage, SupabaseStorage, reset_storage_for_tests
from app.services.permission_gate import _category_for
from app.tools import document_generation

MAGIC = {
    "docx": b"PK",
    "xlsx": b"PK",
    "pdf": b"%PDF",
    "csv": None,
    "md": None,
    "txt": None,
    "html": b"<!DOCTYPE html>",
    "pptx": b"PK",
}

MARKDOWN = (
    "# Quarterly Report\n\n"
    "Revenue grew **12%** year over year.\n\n"
    "- North: up\n- South: flat\n\n"
    "| Region | Growth |\n|---|---|\n| North | 12% |\n| South | 0% |\n"
)

SHEET = {"columns": ["Region", "Growth"], "rows": [["North", 12], ["South", 0]]}


@pytest.fixture(autouse=True)
def local_bucket(tmp_path, monkeypatch):
    """Every test writes into its own bucket directory, never the app's."""
    monkeypatch.setattr(settings, "generated_file_directory", str(tmp_path / "files"))
    reset_storage_for_tests()
    yield tmp_path / "files"
    reset_storage_for_tests()


@pytest.fixture
def owner() -> tuple[str, str]:
    initialize_database()
    user_id = f"gen-{uuid4()}"
    with SessionLocal() as db:
        db.add(User(id=user_id))
        db.commit()
        conversation = Conversation(id=str(uuid4()), user_id=user_id, title="Owner thread")
        db.add(conversation)
        db.commit()
        return user_id, conversation.id


@pytest.fixture
def stranger() -> tuple[str, str]:
    user_id = f"gen-{uuid4()}"
    with SessionLocal() as db:
        db.add(User(id=user_id))
        db.commit()
        conversation = Conversation(id=str(uuid4()), user_id=user_id, title="Stranger thread")
        db.add(conversation)
        db.commit()
        return user_id, conversation.id


def _ctx(user_id: str, conversation_id: str) -> ToolContext:
    return ToolContext(
        user_id=user_id,
        conversation_id=conversation_id,
        turn_id=f"turn-{uuid4()}",
        permission_level=2,
    )


def _generate(user_id: str, conversation_id: str, format: str, title: str = "Report", content=None) -> dict:
    if content is None:
        content = SHEET if format in {"xlsx", "csv"} else MARKDOWN
    result = asyncio.run(
        document_generation.run(
            {"format": format, "title": title, "content": content},
            _ctx(user_id, conversation_id),
        )
    )
    assert result.ok, f"{format} generation failed: {result.error}"
    return result.data


# --- generation -------------------------------------------------------------


@pytest.mark.parametrize("format", ["docx", "pdf", "xlsx", "csv", "md", "txt", "html", "pptx"])
def test_each_format_produces_a_real_file(format, owner, local_bucket):
    user_id, conversation_id = owner
    data = _generate(user_id, conversation_id, format)

    assert data["format"] == format
    assert data["filename"].endswith(f".{format}")
    assert data["conversation_id"] == conversation_id
    assert data["size"] > 0
    assert data["id"] and data["created_at"]

    stored = local_bucket / data["storage_path"]
    assert stored.exists(), "the record must point at bytes that exist"
    raw = stored.read_bytes()
    assert len(raw) == data["size"]
    if MAGIC[format]:
        assert raw.startswith(MAGIC[format]), f"{format} does not start with its magic bytes"

    with SessionLocal() as db:
        row = db.get(GeneratedFile, data["id"])
        assert row is not None and row.user_id == user_id
        assert row.size_bytes == data["size"]


def test_generation_rejects_bad_input(owner):
    user_id, conversation_id = owner
    for args in (
        {"format": "exe", "title": "x", "content": MARKDOWN},
        {"format": "md", "title": "", "content": MARKDOWN},
        {"format": "md", "title": "x", "content": "   "},
        {"format": "xlsx", "title": "x", "content": {"rows": []}},
    ):
        result = asyncio.run(
            document_generation.run(dict(args), _ctx(user_id, conversation_id))
        )
        assert not result.ok, f"{args} should have been rejected"
        assert result.error.code == "INVALID_ARGS"


# --- serving ----------------------------------------------------------------


def test_metadata_download_and_preview(owner, local_bucket):
    user_id, conversation_id = owner
    record = _generate(user_id, conversation_id, "md", title="Notes")
    client = TestClient(app)

    meta = client.get("/api/files/" + record["id"], params={"user_id": user_id})
    assert meta.status_code == 200
    assert meta.json()["filename"] == "Notes.md"
    assert meta.json()["conversation_id"] == conversation_id

    download = client.get(f"/api/files/{record['id']}/download", params={"user_id": user_id})
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("text/markdown")
    disposition = download.headers["content-disposition"]
    assert disposition.startswith("attachment;") and "Notes.md" in disposition
    # Byte-identical to what was generated: the preview must never disagree
    # with the download.
    assert download.content == (local_bucket / record["storage_path"]).read_bytes()
    assert download.content.decode("utf-8") == MARKDOWN

    preview = client.get(f"/api/files/{record['id']}/preview", params={"user_id": user_id})
    assert preview.status_code == 200
    body = preview.json()
    assert body["preview"] == "inline"
    assert body["text"] == MARKDOWN
    assert "storage_path" not in body, "the preview must not leak the bucket key"


def test_binary_preview_defers_to_bytes(owner):
    user_id, conversation_id = owner
    record = _generate(user_id, conversation_id, "xlsx", title="Budget")
    client = TestClient(app)
    preview = client.get(f"/api/files/{record['id']}/preview", params={"user_id": user_id})
    assert preview.status_code == 200
    assert preview.json()["preview"] == "bytes"
    assert "text" not in preview.json()
    # ...and the bytes it points at are a real workbook.
    download = client.get(f"/api/files/{record['id']}/download", params={"user_id": user_id})
    assert download.status_code == 200
    assert download.content.startswith(b"PK")


def test_conversation_and_user_listings_are_scoped(owner, stranger):
    user_id, conversation_id = owner
    other_user, other_conversation = stranger
    record = _generate(user_id, conversation_id, "csv", title="Export")
    client = TestClient(app)

    mine = client.get(
        f"/api/conversations/{conversation_id}/files", params={"user_id": user_id}
    )
    assert mine.status_code == 200
    assert [entry["id"] for entry in mine.json()["files"]] == [record["id"]]

    # Another user's conversation: not found, not empty-but-readable.
    theirs = client.get(
        f"/api/conversations/{conversation_id}/files", params={"user_id": other_user}
    )
    assert theirs.status_code == 404

    their_own = client.get(
        f"/api/conversations/{other_conversation}/files", params={"user_id": other_user}
    )
    assert their_own.status_code == 200 and their_own.json()["files"] == []

    assert client.get("/api/files", params={"user_id": user_id}).json()["files"]
    assert client.get("/api/files", params={"user_id": other_user}).json()["files"] == []


def test_a_file_is_invisible_to_everyone_else(owner, stranger, local_bucket):
    user_id, conversation_id = owner
    other_user, _ = stranger
    record = _generate(user_id, conversation_id, "pdf", title="Secret")
    client = TestClient(app)

    for path in (
        f"/api/files/{record['id']}",
        f"/api/files/{record['id']}/download",
        f"/api/files/{record['id']}/preview",
    ):
        assert client.get(path, params={"user_id": other_user}).status_code == 404, path

    delete = client.delete(f"/api/files/{record['id']}", params={"user_id": other_user})
    assert delete.status_code == 404

    # The stranger's failed delete must not have touched anything.
    assert (local_bucket / record["storage_path"]).exists()
    assert client.get(f"/api/files/{record['id']}", params={"user_id": user_id}).status_code == 200


def test_delete_removes_the_object_and_the_row(owner, local_bucket):
    user_id, conversation_id = owner
    record = _generate(user_id, conversation_id, "docx", title="Draft")
    stored = local_bucket / record["storage_path"]
    assert stored.exists()
    client = TestClient(app)

    deleted = client.delete(f"/api/files/{record['id']}", params={"user_id": user_id})
    assert deleted.status_code == 200 and deleted.json()["deleted"] is True
    assert not stored.exists(), "the bytes must go, not just the row"
    with SessionLocal() as db:
        assert db.get(GeneratedFile, record["id"]) is None
    assert client.get(f"/api/files/{record['id']}", params={"user_id": user_id}).status_code == 404
    # Empty per-conversation folders are pruned with the file.
    assert not stored.parent.exists()


# --- storage backends -------------------------------------------------------


def test_key_is_scoped_and_sanitised():
    from app.services.file_storage import StorageError, _safe_key, build_key

    key = build_key("../../etc", "conv/../x", "file", "pdf")
    parts = key.split("/")
    assert len(parts) == 3, "a key is always user/conversation/file.ext"
    assert all(part not in {"", ".", ".."} for part in parts)
    assert "/" not in parts[0], "the user component can never hold a separator"
    assert key.endswith(".pdf")
    # Traversal is refused at the door, not merely sanitised away.
    for bad in ("../escape", "/abs/path", "a//b", "", "a/../../b"):
        with pytest.raises(StorageError):
            _safe_key(bad)


def test_supabase_backend_uses_private_bucket_and_service_key(tmp_path):
    """Storage REST calls carry the service-role key and hit bucket paths."""
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        if request.method == "GET":
            return httpx.Response(200, content=b"bytes")
        return httpx.Response(200, json={"Key": "ok"})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    storage = SupabaseStorage(
        "https://proj.supabase.co", "service-role-secret", "pentagon-files", client=client
    )

    async def roundtrip():
        await storage.put("user/conv/file.md", b"# hi", "text/markdown")
        body = await storage.get("user/conv/file.md")
        await storage.delete("user/conv/file.md")
        return body

    assert asyncio.run(roundtrip()) == b"bytes"
    methods = [request.method for request in seen]
    assert methods[0] == "POST"  # ensure bucket
    assert methods[1] == "POST"  # upload
    upload = seen[1]
    assert "/storage/v1/object/pentagon-files/user/conv/file.md" in str(upload.url)
    assert upload.headers["apikey"] == "service-role-secret"
    assert upload.headers["Authorization"] == "Bearer service-role-secret"
    assert "GET" in methods and "DELETE" in methods
    assert seen[-1].method == "DELETE"
    asyncio.run(client.aclose())


def test_storage_choice_follows_configuration(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "generated_file_directory", str(tmp_path / "local"))
    reset_storage_for_tests()
    assert isinstance(file_storage.get_storage(), LocalStorage)

    from pydantic import SecretStr

    monkeypatch.setattr(settings, "supabase_url", "https://proj.supabase.co")
    monkeypatch.setattr(settings, "supabase_service_role_key", SecretStr("svc"))
    reset_storage_for_tests()
    try:
        storage = file_storage.get_storage()
        assert isinstance(storage, SupabaseStorage)
        assert storage.bucket == "pentagon-files"
    finally:
        reset_storage_for_tests()


# --- the gate ---------------------------------------------------------------


def test_document_generation_is_classified_as_a_file_write():
    """The tool must be recognised by the gate as a local filesystem write."""
    assert document_generation.TOOL_SPEC.risk_category == "local_filesystem_write"
    assert document_generation.TOOL_SPEC.tier == "write"
    assert _category_for("document_generation", {"format": "pdf"}) == "local_filesystem_write"


def _scripted_model(responses):
    class Model:
        def __init__(self):
            self.calls = 0

        def bind_tools(self, tools):
            return self

        async def ainvoke(self, messages, config=None):
            self.calls += 1
            return responses.pop(0) if responses else AIMessage(content="done")

    return Model()


def _generation_call() -> AIMessage:
    return AIMessage(
        content="",
        tool_calls=[
            {
                "name": "document_generation",
                "args": {"format": "md", "title": "Gate Check", "content": "# Hi\n\nBody text."},
                "id": "call_1",
            }
        ],
    )


def _graph(model, user_id, monkeypatch):
    monkeypatch.setattr(chat_graph, "make_chat_model", lambda *_a, **_k: model)
    monkeypatch.setattr(chat_graph, "has_documents", lambda _c: False)
    return chat_graph.build_chat_graph(
        "unused-key",
        "test-model",
        user_id=user_id,
        use_web_search=False,
        command_tool_enabled=False,
        checkpointer=InMemorySaver(),
    )


def _state(user_id: str, conversation_id: str):
    return chat_graph.initial_chat_state(
        user_id=user_id,
        conversation_id=conversation_id,
        message="make me a file",
        history=[],
        use_web_search=False,
        command_tool_enabled=False,
    )


def test_generating_a_file_pauses_at_the_gate_and_runs_after_approval(
    owner, local_bucket, monkeypatch
):
    """local_filesystem_write defaults to always_ask: nothing is written
    until the user approves, then the approved call really writes."""
    user_id, conversation_id = owner
    model = _scripted_model([_generation_call(), AIMessage(content="Done.")])
    graph = _graph(model, user_id, monkeypatch)
    config = {"configurable": {"thread_id": conversation_id}}

    asyncio.run(graph.ainvoke(_state(user_id, conversation_id), config=config))
    snap = asyncio.run(graph.aget_state(config))
    cards = [interrupt.value for interrupt in (snap.interrupts or ())]
    assert cards, "file generation must stop for approval at the default level"
    assert cards[0]["type"] == "approval_required"
    assert cards[0]["category"] == "local_filesystem_write"

    with SessionLocal() as db:
        rows = db.query(GeneratedFile).filter(GeneratedFile.user_id == user_id).all()
    assert rows == [], "a file was written before the user approved"
    assert not any(local_bucket.rglob("*.*")), "storage was written before approval"

    asyncio.run(
        graph.ainvoke(
            Command(resume={"decision": "approved", "approved_this_session": True}),
            config=config,
        )
    )
    with SessionLocal() as db:
        rows = db.query(GeneratedFile).filter(GeneratedFile.user_id == user_id).all()
    assert len(rows) == 1, "the approved generation must produce a file"
    assert (local_bucket / rows[0].storage_path).exists()

    gate_rows = [row for row in _audit_rows(user_id) if row.tool_name == "document_generation"]
    assert gate_rows, "the audit trail must record the generation attempt"
    assert gate_rows[-1].decision == "user_approved"
    assert gate_rows[-1].category == "local_filesystem_write"


def test_generating_a_file_is_blocked_when_never_allowed(owner, monkeypatch, local_bucket):
    user_id, conversation_id = owner
    with SessionLocal() as db:
        db.add(
            AutonomySetting(
                user_id=user_id, category="local_filesystem_write", level="never_allow"
            )
        )
        db.commit()
    model = _scripted_model([_generation_call(), AIMessage(content="Understood.")])
    graph = _graph(model, user_id, monkeypatch)
    config = {"configurable": {"thread_id": conversation_id}}

    asyncio.run(graph.ainvoke(_state(user_id, conversation_id), config=config))
    with SessionLocal() as db:
        rows = db.query(GeneratedFile).filter(GeneratedFile.user_id == user_id).all()
    assert rows == [], "never_allow must block the write outright"
    assert not any(local_bucket.rglob("*.*"))


def test_generating_a_file_runs_unattended_when_auto_approved(owner, monkeypatch, local_bucket):
    user_id, conversation_id = owner
    with SessionLocal() as db:
        db.add(
            AutonomySetting(
                user_id=user_id, category="local_filesystem_write", level="auto_approve"
            )
        )
        db.commit()
    model = _scripted_model([_generation_call(), AIMessage(content="Done.")])
    graph = _graph(model, user_id, monkeypatch)
    config = {"configurable": {"thread_id": conversation_id}}

    result = asyncio.run(graph.ainvoke(_state(user_id, conversation_id), config=config))
    snap = asyncio.run(graph.aget_state(config))
    assert not (snap.interrupts or ()), "auto_approve must not interrupt"
    with SessionLocal() as db:
        rows = db.query(GeneratedFile).filter(GeneratedFile.user_id == user_id).all()
    assert len(rows) == 1
    assert (local_bucket / rows[0].storage_path).exists()
    assert result is not None


def _audit_rows(user_id: str):
    with SessionLocal() as db:
        return list(
            db.query(ActionAuditLog)
            .filter(ActionAuditLog.user_id == user_id)
            .order_by(ActionAuditLog.timestamp)
            .all()
        )


def test_multipage_document_generation_with_numbering_and_breaks(owner, local_bucket):
    import io
    from pypdf import PdfReader
    import openpyxl
    from pptx import Presentation
    from docx import Document as WordDocument

    user_id, conversation_id = owner

    # 1. Multi-page PDF with page break directives
    multipage_md = (
        "# Executive Summary\n\n"
        "This is the first page of the report.\n\n"
        "---\n\n"
        "# Financial Performance\n\n"
        "This is the second page of the report.\n\n"
        "---\n\n"
        "# Conclusion\n\n"
        "This is the third page of the report.\n"
    )
    pdf_file = _generate(user_id, conversation_id, "pdf", title="Annual Report", content=multipage_md)
    pdf_bytes = (local_bucket / pdf_file["storage_path"]).read_bytes()
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) == 3, f"Expected 3 pages, got {len(reader.pages)}"
    p1_text = reader.pages[0].extract_text()
    p2_text = reader.pages[1].extract_text()
    p3_text = reader.pages[2].extract_text()
    assert "Page 1 of 3" in p1_text
    assert "Page 2 of 3" in p2_text
    assert "Page 3 of 3" in p3_text
    assert "Pentagon Generated Document" in p1_text

    # 2. Multi-page DOCX with page breaks
    docx_file = _generate(user_id, conversation_id, "docx", title="Word Report", content=multipage_md)
    docx_bytes = (local_bucket / docx_file["storage_path"]).read_bytes()
    w_doc = WordDocument(io.BytesIO(docx_bytes))
    assert len(w_doc.paragraphs) >= 3

    # 3. Multi-slide PPTX
    pptx_file = _generate(user_id, conversation_id, "pptx", title="Roadmap Presentation", content=multipage_md)
    pptx_bytes = (local_bucket / pptx_file["storage_path"]).read_bytes()
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) >= 3, f"Expected at least 3 slides, got {len(prs.slides)}"

    # 4. Multi-sheet XLSX
    sheet_data = {
        "sheets": [
            {"name": "2025 Revenue", "columns": ["Quarter", "Revenue"], "rows": [["Q1", 100], ["Q2", 150]]},
            {"name": "2026 Projections", "columns": ["Quarter", "Forecast"], "rows": [["Q1", 200], ["Q2", 250]]},
        ]
    }
    xlsx_file = _generate(user_id, conversation_id, "xlsx", title="Financial Model", content=sheet_data)
    xlsx_bytes = (local_bucket / xlsx_file["storage_path"]).read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["2025 Revenue", "2026 Projections"]


def test_robust_doc_input_formats_and_aliases(owner, local_bucket):
    user_id, conversation_id = owner

    # 1. Dict with "text" or "content" key
    doc1 = _generate(
        user_id,
        conversation_id,
        "docx",
        title="Alias Test",
        content={"text": "# Title with text key\n\nContent paragraph here."},
    )
    assert doc1["format"] == "docx"
    assert (local_bucket / doc1["storage_path"]).exists()

    # 2. Markdown table string directly for xlsx
    md_table = "| Metric | Value |\n|---|---|\n| Growth | 25% |\n| Retention | 90% |\n"
    doc2 = _generate(
        user_id,
        conversation_id,
        "xlsx",
        title="Metrics Sheet",
        content=md_table,
    )
    assert doc2["format"] == "xlsx"
    assert (local_bucket / doc2["storage_path"]).exists()

    # 3. Unicode title in PDF
    doc3 = _generate(
        user_id,
        conversation_id,
        "pdf",
        title="🚀 Q3 Executive Report — Global (2026)",
        content="# Overview\n\nThis is a unicode test.",
    )
    assert doc3["format"] == "pdf"
    assert (local_bucket / doc3["storage_path"]).exists()


